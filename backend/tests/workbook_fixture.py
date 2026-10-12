from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook


def realistic_workbook(*, malformed: bool = False, months: tuple[str, ...] = ("August", "September")) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.modified = datetime(2026, 9, 30)
    for index, month in enumerate(months):
        sheet = workbook.create_sheet(month)
        sheet.append(["Income"])
        sheet.append(["Amount", "Description / Type", "Income Class"])
        sheet.append([4200 + index * 100, "Primary paycheck", "Paycheck"])
        sheet.append([250, "Freelance project", "Side Income"])
        sheet.append([])
        sheet.append([])
        sheet.append(["Expenses"])
        sheet.append(["Amount", "Description / Type", "Category", "Expense Tag"])
        sheet.append([1500, "Rent", "Housing", ""])
        sheet.append([85, "Client dinner", "Work Meals", "Work/Reimbursable"])
        sheet.append([120, "Weekend trip", "Travel", "Personal Travel"])
        sheet.append([300, "New desk", "Shopping & Home" if not malformed else "", "One-time"])
        sheet.append([160, "Groceries", "Groceries", ""])
        sheet.append([])
        sheet.append([])
        sheet.append(["Wealth Flow"])
        sheet.append(["Amount", "Destination", "Flow Class"])
        sheet.append([500, "401(k)", "Retirement Investment"])
        sheet.append([300, "Roth IRA", "Investment"])
        sheet.append([1000, "Chase -> HYSA", "Internal Transfer"])
    net_worth = workbook.create_sheet("Networth")
    net_worth.append(["Snapshot Date", date(2026, 9, 30)])
    net_worth.append([])
    net_worth.append(["Account", "Account Type", "Value"])
    net_worth.append(["Checking", "Cash", 4500])
    net_worth.append(["Citi HYSA", "Cash", 12000])
    net_worth.append(["Vanguard", "Taxable Brokerage", 18000])
    net_worth.append(["Roth IRA", "Roth IRA", 9000])
    net_worth.append(["401(k)", "401(k)", 22000])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
