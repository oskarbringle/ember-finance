from __future__ import annotations

import re
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from io import BytesIO
from typing import Any

from openpyxl import load_workbook


MONTHS = {
    name.lower(): index
    for index, name in enumerate(
        (
            "January",
            "February",
            "March",
            "April",
            "May",
            "June",
            "July",
            "August",
            "September",
            "October",
            "November",
            "December",
        ),
        1,
    )
}
MONTHS.update({name[:3]: number for name, number in list(MONTHS.items())})
KNOWN_REFERENCE_SHEETS = {"categories"}


def normalized(value: Any) -> str:
    text = str(value or "").strip().lower()
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def parse_amount(value: Any) -> float:
    if isinstance(value, (int, float)):
        return abs(float(value))
    text = str(value or "").strip()
    if not text:
        raise ValueError("amount is blank")
    negative = text.startswith("(") and text.endswith(")")
    cleaned = re.sub(r"[^0-9.\-]", "", text)
    if not cleaned or cleaned in {"-", ".", "-."}:
        raise ValueError(f"'{text}' is not a number")
    amount = float(cleaned)
    return abs(-amount if negative else amount)


def parse_date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    embedded_date = re.search(
        r"(?:\d{1,2}/\d{1,2}/(?:\d{4}|\d{2})\b|[A-Za-z]+\s+\d{1,2},\s+\d{4})",
        text,
    )
    if embedded_date:
        text = embedded_date.group(0)
    for format_string in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%B %d, %Y"):
        try:
            return datetime.strptime(text, format_string).date()
        except ValueError:
            continue
    raise ValueError(f"'{text}' is not a recognized date")


@dataclass(frozen=True)
class ImportIssue:
    sheet: str
    row: int | None
    message: str
    severity: str = "error"


@dataclass(frozen=True)
class ParsedIncome:
    month: str
    amount: float
    description: str
    income_class: str


@dataclass(frozen=True)
class ParsedExpense:
    month: str
    amount: float
    description: str
    category: str
    expense_tag: str = ""


@dataclass(frozen=True)
class ParsedWealthFlow:
    month: str
    amount: float
    destination: str
    flow_class: str


@dataclass(frozen=True)
class ParsedNetWorth:
    snapshot_date: date
    account: str
    account_type: str
    value: float


@dataclass
class ParsedWorkbook:
    workbook_name: str
    detected_months: list[str] = field(default_factory=list)
    income: list[ParsedIncome] = field(default_factory=list)
    expenses: list[ParsedExpense] = field(default_factory=list)
    wealth_flows: list[ParsedWealthFlow] = field(default_factory=list)
    net_worth: list[ParsedNetWorth] = field(default_factory=list)
    issues: list[ImportIssue] = field(default_factory=list)

    def preview(self, token: str) -> dict:
        return {
            "token": token,
            "workbook_name": self.workbook_name,
            "detected_months": self.detected_months,
            "income_entries": len(self.income),
            "expense_entries": len(self.expenses),
            "wealth_flow_entries": len(self.wealth_flows),
            "net_worth_entries": len(self.net_worth),
            "net_worth_dates": sorted(
                {row.snapshot_date.isoformat() for row in self.net_worth}
            ),
            "issues": [asdict(issue) for issue in self.issues],
            "can_import": not any(issue.severity == "error" for issue in self.issues)
            and bool(self.detected_months),
        }


HEADER_SCHEMAS = {
    "income": {
        "amount": {"amount", "income", "income amount"},
        "description": {"description", "type", "description type", "income type"},
        "income_class": {"income class", "class"},
    },
    "expenses": {
        "amount": {"amount", "expense", "expenses", "expense amount"},
        "description": {"description", "type", "description type", "expense type"},
        "category": {"category", "expense category"},
        "expense_tag": {"expense tag", "tag", "tags"},
    },
    "wealth_flows": {
        "amount": {"amount", "wealth flow", "flow amount"},
        "destination": {"destination", "account", "flow destination"},
        "flow_class": {"flow class", "class", "wealth flow class"},
    },
}
REQUIRED_FIELDS = {
    "income": {"amount", "description", "income_class"},
    "expenses": {"amount", "description", "category"},
    "wealth_flows": {"amount", "destination", "flow_class"},
}


def _month_from_title(title: str, default_year: int) -> tuple[int, int] | None:
    match = re.fullmatch(r"\s*([A-Za-z]+)(?:[\s_-]+(20\d{2}))?\s*", title)
    if not match:
        return None
    month_number = MONTHS.get(match.group(1).lower())
    if not month_number:
        return None
    return int(match.group(2) or default_year), month_number


def _month_keys(workbook) -> dict[str, str]:
    modified = workbook.properties.modified
    default_year = modified.year if modified else datetime.now().year
    result: dict[str, str] = {}
    prior_month = 0
    year = default_year
    for title in workbook.sheetnames:
        parsed = _month_from_title(title, year)
        if not parsed:
            continue
        explicit_year, month_number = parsed
        if re.search(r"20\d{2}", title):
            year = explicit_year
        elif prior_month and month_number < prior_month:
            year += 1
        result[title] = f"{year:04d}-{month_number:02d}"
        prior_month = month_number
    return result


def _find_headers(sheet) -> list[tuple[str, int, dict[str, int]]]:
    found: list[tuple[str, int, dict[str, int]]] = []
    for row_number in range(1, sheet.max_row + 1):
        values = {
            column: normalized(sheet.cell(row_number, column).value)
            for column in range(1, sheet.max_column + 1)
        }
        for start in range(1, sheet.max_column + 1):
            window = range(start, min(sheet.max_column, start + 6) + 1)
            for section, aliases in HEADER_SCHEMAS.items():
                if values[start] not in aliases["amount"]:
                    continue
                mapping = {"amount": start}
                for field_name, names in aliases.items():
                    if field_name == "amount":
                        continue
                    match = next(
                        (column for column in window if values[column] in names),
                        None,
                    )
                    if match:
                        mapping[field_name] = match
                if REQUIRED_FIELDS[section].issubset(mapping):
                    candidate = (section, row_number, mapping)
                    if candidate not in found:
                        found.append(candidate)
    return found


def _read_section(
    sheet,
    section: str,
    header_row: int,
    columns: dict[str, int],
    month: str,
    parsed: ParsedWorkbook,
) -> None:
    primary_fields = {
        "income": ("amount", "description"),
        "expenses": ("amount", "description"),
        "wealth_flows": ("amount", "destination"),
    }[section]
    primary_columns = tuple(columns[field_name] for field_name in primary_fields)
    last_data_row = max(
        (
            row_number
            for row_number in range(header_row + 1, sheet.max_row + 1)
            if any(
                sheet.cell(row_number, column).value not in (None, "")
                for column in primary_columns
            )
        ),
        default=header_row,
    )
    for row_number in range(header_row + 1, last_data_row + 1):
        values = {
            field_name: sheet.cell(row_number, column).value
            for field_name, column in columns.items()
        }
        if not any(values.get(field_name) not in (None, "") for field_name in primary_fields):
            continue
        if normalized(values.get("amount")) in HEADER_SCHEMAS[section]["amount"]:
            break
        missing = [
            field_name
            for field_name in REQUIRED_FIELDS[section]
            if values.get(field_name) in (None, "")
        ]
        if missing:
            parsed.issues.append(
                ImportIssue(sheet.title, row_number, f"Missing {', '.join(sorted(missing))} in {section.replace('_', ' ')} row")
            )
            continue
        try:
            amount = parse_amount(values["amount"])
        except (TypeError, ValueError) as exc:
            parsed.issues.append(ImportIssue(sheet.title, row_number, f"Invalid amount: {exc}"))
            continue
        if section == "income":
            parsed.income.append(
                ParsedIncome(month, amount, str(values["description"]).strip(), str(values["income_class"]).strip())
            )
        elif section == "expenses":
            parsed.expenses.append(
                ParsedExpense(
                    month,
                    amount,
                    str(values["description"]).strip(),
                    str(values["category"]).strip(),
                    str(values.get("expense_tag") or "").strip(),
                )
            )
        else:
            parsed.wealth_flows.append(
                ParsedWealthFlow(month, amount, str(values["destination"]).strip(), str(values["flow_class"]).strip())
            )


def _shared_snapshot_date(sheet) -> date | None:
    for row in sheet.iter_rows():
        for cell in row:
            try:
                if any(token in normalized(cell.value) for token in ("snapshot date", "as of")):
                    return parse_date(cell.value)
            except ValueError:
                pass
            if normalized(cell.value) not in {"snapshot date", "as of", "date"}:
                continue
            for candidate in (
                sheet.cell(cell.row, cell.column + 1).value,
                sheet.cell(cell.row + 1, cell.column).value,
            ):
                if candidate not in (None, ""):
                    try:
                        return parse_date(candidate)
                    except ValueError:
                        pass
    return None


def _investment_account_type(account: str) -> str:
    account_name = normalized(account)
    if "roth" in account_name:
        return "Roth IRA"
    if "hsa" in account_name:
        return "HSA"
    if "401" in account_name:
        return "401(k)"
    return "Taxable Brokerage"


def _read_net_worth_blocks(sheet, parsed: ParsedWorkbook, snapshot_date: date | None) -> bool:
    headers = {
        normalized(sheet.cell(1, column).value): column
        for column in range(1, sheet.max_column + 1)
    }
    if "savings" not in headers or "investments" not in headers:
        return False
    if not snapshot_date:
        parsed.issues.append(
            ImportIssue(sheet.title, 1, "Networth tab needs an 'as of' date")
        )
        return True
    for account_column, fixed_type in (
        (headers["savings"], "Cash"),
        (headers["investments"], None),
    ):
        value_column = account_column + 1
        for row_number in range(2, sheet.max_row + 1):
            account = sheet.cell(row_number, account_column).value
            value = sheet.cell(row_number, value_column).value
            if account in (None, "") and value in (None, ""):
                continue
            if normalized(account).startswith("total"):
                break
            if account in (None, "") or value in (None, ""):
                parsed.issues.append(
                    ImportIssue(
                        sheet.title,
                        row_number,
                        "Net-worth row needs an account and value",
                    )
                )
                continue
            try:
                account_name = str(account).strip()
                account_type = fixed_type or _investment_account_type(account_name)
                parsed.net_worth.append(
                    ParsedNetWorth(
                        snapshot_date,
                        account_name,
                        account_type,
                        parse_amount(value),
                    )
                )
            except (TypeError, ValueError) as exc:
                parsed.issues.append(
                    ImportIssue(
                        sheet.title,
                        row_number,
                        f"Invalid net-worth row: {exc}",
                    )
                )
    return True


def _read_net_worth(sheet, parsed: ParsedWorkbook) -> None:
    shared_date = _shared_snapshot_date(sheet)
    if _read_net_worth_blocks(sheet, parsed, shared_date):
        return
    for row_number in range(1, sheet.max_row + 1):
        values = {
            column: normalized(sheet.cell(row_number, column).value)
            for column in range(1, sheet.max_column + 1)
        }
        account_columns = [column for column, value in values.items() if value in {"account", "account name"}]
        for start in account_columns:
            window = range(start, min(sheet.max_column, start + 6) + 1)
            mapping: dict[str, int] = {"account": start}
            aliases = {
                "account_type": {"account type", "type", "group"},
                "value": {"value", "balance", "current balance", "amount"},
                "date": {"date", "snapshot date", "as of"},
            }
            for field_name, names in aliases.items():
                match = next((column for column in window if values[column] in names), None)
                if match:
                    mapping[field_name] = match
            if not {"account", "account_type", "value"}.issubset(mapping):
                continue
            blank_rows = 0
            for data_row in range(row_number + 1, sheet.max_row + 1):
                account = sheet.cell(data_row, mapping["account"]).value
                account_type = sheet.cell(data_row, mapping["account_type"]).value
                value = sheet.cell(data_row, mapping["value"]).value
                if account in (None, "") and value in (None, ""):
                    blank_rows += 1
                    if blank_rows >= 2:
                        break
                    continue
                blank_rows = 0
                if account in (None, "") or account_type in (None, "") or value in (None, ""):
                    parsed.issues.append(ImportIssue(sheet.title, data_row, "Net-worth row needs account, account type, and value"))
                    continue
                try:
                    snapshot_date = parse_date(sheet.cell(data_row, mapping["date"]).value) if "date" in mapping else shared_date
                    if not snapshot_date:
                        raise ValueError("add a Snapshot Date above the table or a Date column")
                    parsed.net_worth.append(
                        ParsedNetWorth(snapshot_date, str(account).strip(), str(account_type).strip(), parse_amount(value))
                    )
                except (TypeError, ValueError) as exc:
                    parsed.issues.append(ImportIssue(sheet.title, data_row, f"Invalid net-worth row: {exc}"))
            return
    parsed.issues.append(ImportIssue(sheet.title, None, "Could not find Account, Account Type, and Value headers"))


def parse_workbook(content: bytes, filename: str) -> ParsedWorkbook:
    parsed = ParsedWorkbook(filename)
    try:
        workbook = load_workbook(BytesIO(content), data_only=True, read_only=False)
    except Exception as exc:
        parsed.issues.append(ImportIssue("Workbook", None, f"Could not open workbook: {exc}"))
        return parsed
    month_keys = _month_keys(workbook)
    for title, month in month_keys.items():
        sheet = workbook[title]
        headers = _find_headers(sheet)
        if not headers:
            parsed.issues.append(ImportIssue(title, None, "Month tab has no recognizable Income, Expenses, or Wealth Flow headers"))
            continue
        parsed.detected_months.append(month)
        seen: set[tuple[str, int, tuple[int, ...]]] = set()
        for section, header_row, columns in headers:
            key = (section, header_row, tuple(sorted(columns.values())))
            if key in seen:
                continue
            seen.add(key)
            _read_section(sheet, section, header_row, columns, month, parsed)
    net_worth_sheet = next(
        (sheet for sheet in workbook.worksheets if normalized(sheet.title) in {"networth", "net worth"}),
        None,
    )
    if net_worth_sheet:
        _read_net_worth(net_worth_sheet, parsed)
    else:
        parsed.issues.append(ImportIssue("Workbook", None, "No Networth tab detected", "warning"))
    ignored = [
        title
        for title in workbook.sheetnames
        if title not in month_keys
        and (not net_worth_sheet or title != net_worth_sheet.title)
        and normalized(title) not in KNOWN_REFERENCE_SHEETS
    ]
    for title in ignored:
        parsed.issues.append(ImportIssue(title, None, "Sheet was not recognized and was not imported", "warning"))
    parsed.detected_months = sorted(set(parsed.detected_months))
    return parsed
